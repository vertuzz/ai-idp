"""Bounded, inert workbook decoding with sparse OOXML cells and stable coordinates."""

from __future__ import annotations

import io
import json
import posixpath
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any
from zipfile import ZipFile

from defusedxml import ElementTree as XML
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from ai_idp.adapters.outbound.documents.formulas import (
    FormulaCell,
    FormulaEvaluator,
    UnsupportedFormula,
)
from ai_idp.application.ports.protocols import DocumentBlobStore
from ai_idp.domain.enums import UnitKind
from ai_idp.domain.errors import (
    DocumentLimitExceededError,
    EncryptedDocumentError,
    MalformedDocumentError,
)
from ai_idp.domain.media_types import SPREADSHEET_TYPES, XLSM_MIME, XLSX_MIME
from ai_idp.domain.models import Document, DocumentUnit
from ai_idp.settings import Settings

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def formatted_value(value: str | None, kind: str, number_format: str) -> str | None:
    """Render only unambiguous simple masks; retain raw values and all format codes.

    This is mechanical representation, not formula evaluation or an Office renderer.
    Complex/localized formats remain explicit metadata for the model.
    """
    if kind == "b" and value is not None:
        if value not in {"0", "1", "true", "false"}:
            raise MalformedDocumentError("Invalid boolean workbook value.")
        return "TRUE" if value in {"1", "true"} else "FALSE"
    if kind != "n" or value is None:
        return value
    try:
        number = Decimal(value)
        if not number.is_finite() or abs(number.adjusted()) > 100 or len(number_format) > 100:
            return value
        if re.fullmatch(r"0{2,}", number_format) and number >= 0 and number == number.to_integral():
            return str(int(number)).zfill(len(number_format))
        percent = re.fullmatch(r"0(?:\.(0+))?%", number_format)
        if percent:
            precision = len(percent.group(1) or "")
            if precision > 20:
                return value
            return f"{number * 100:.{precision}f}%"
    except (InvalidOperation, ValueError):
        pass
    return value


def _part(base: str, target: str) -> str:
    name = posixpath.normpath(posixpath.join(posixpath.dirname(base), target))
    if target.startswith("/"):
        name = target.lstrip("/")
    if name.startswith("../") or ":" in name or "\\" in name:
        raise MalformedDocumentError("Invalid workbook relationship.")
    return name


class SpreadsheetDocumentReader:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def supports(self, media_type: str) -> bool:
        return media_type in SPREADSHEET_TYPES

    def normalize(
        self, document: Document, raw_bytes: bytes, blob_store: DocumentBlobStore
    ) -> list[DocumentUnit]:
        self.units: list[DocumentUnit] = []
        self.count = self.characters = 0
        self.document, self.store = document, blob_store
        if document.media_type in {XLSX_MIME, XLSM_MIME}:
            self._ooxml(raw_bytes)
        else:
            self._calamine(raw_bytes)
        if not self.units:
            self._emit(1, "Workbook", "visible", [], ["Workbook contains no readable cells."])
        return self.units

    def _emit(
        self,
        sheet_index: int,
        name: str,
        state: str,
        cells: list[dict[str, Any]],
        warnings: list[str],
        images: list[str] | None = None,
    ) -> None:
        if len(self.units) >= self.settings.max_document_units:
            raise DocumentLimitExceededError("Workbook source-unit limit exceeded.")
        uid = f"sheet_{sheet_index:04d}_range_{len(self.units) + 1:06d}"
        header = f"Sheet: {name} (index {sheet_index}, state {state})\n"
        lines = [header]
        spans = []
        offset = len(header)
        for cell in cells:
            line = cell["text"] + "\n"
            spans.append({"cell": cell["cell"], "start": offset, "end": offset + len(line) - 1})
            lines.append(line)
            offset += len(line)
        text = "".join(lines)
        self.characters += len(text)
        if self.characters > self.settings.max_document_text_characters:
            raise DocumentLimitExceededError("Workbook decoded-text limit exceeded.")
        base = f"derived/{self.document.document_id}/{uid}"
        locator = {
            "sheet_index": sheet_index,
            "sheet_name": name,
            "cell_range": f"{cells[0]['cell']}:{cells[-1]['cell']}" if cells else None,
        }
        text_ref = self.store.put_text(base + ".txt", text)
        cells_ref = self.store.put_text(base + ".cells.json", json.dumps(cells, ensure_ascii=True))
        self.units.append(
            DocumentUnit(
                unit_id=uid,
                document_id=self.document.document_id,
                ordinal=len(self.units) + 1,
                kind=UnitKind.SPREADSHEET_RANGE,
                source_locator=locator,
                native_text_ref=text_ref,
                image_refs=images or [],
                image_ref=(images or [None])[0],
                metadata={
                    "sheet_state": state,
                    "conversion_warnings": list(warnings),
                    "unavailable_images": list(warnings),
                    "cells_ref": cells_ref,
                    "cell_spans": spans,
                    "cell_calculations": {
                        c["cell"]: c["calculation"] for c in cells if c.get("calculation")
                    },
                    "owned_cells": [c["cell"] for c in cells],
                    "sheet_context": getattr(self, "sheet_context", ""),
                },
            )
        )

    def _chunks(
        self,
        index: int,
        name: str,
        state: str,
        records: Any,
        warnings: list[str],
        images: list[str] | None = None,
    ) -> None:
        chunk: list[dict[str, Any]] = []
        rows: set[int] = set()
        size = 0
        self.sheet_context = ""
        context_rows: set[int] = set()
        for record in records:
            self.count += 1
            if self.count > self.settings.max_spreadsheet_cells:
                raise DocumentLimitExceededError("Workbook populated-cell limit exceeded.")
            row, _ = coordinate_to_tuple(record["cell"])
            if row in context_rows or len(context_rows) < 1:
                context_rows.add(row)
                if len(self.sheet_context) + len(record["text"]) < 2000:
                    self.sheet_context += record["text"] + "\n"
            length = len(record["text"])
            if length > self.settings.max_unit_text_characters:
                raise DocumentLimitExceededError("A workbook cell exceeds the text-unit limit.")
            if chunk and (
                size + length > self.settings.max_unit_text_characters
                or (row not in rows and len(rows) >= self.settings.spreadsheet_rows_per_unit)
            ):
                self._emit(index, name, state, chunk, warnings)
                chunk, rows, size = [], set(), 0
            chunk.append(record)
            rows.add(row)
            size += length + 1
        if (
            chunk
            or images
            or not any(u.source_locator.get("sheet_index") == index for u in self.units)
        ):
            self._emit(index, name, state, chunk, warnings, images)

    def _ooxml(self, raw: bytes) -> None:
        with ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 20_000 or len({i.filename for i in entries}) != len(entries):
                raise DocumentLimitExceededError("Workbook ZIP entry limit or duplicate entry.")
            if sum(i.file_size for i in entries) > self.settings.max_expanded_bytes:
                raise DocumentLimitExceededError("Workbook expanded-size limit exceeded.")
            book = XML.fromstring(archive.read("xl/workbook.xml"))
            if book.tag != NS + "workbook":
                raise MalformedDocumentError("Unsupported workbook XML namespace.")
            relationships = XML.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            rels = {
                r.get("Id"): _part("xl/workbook.xml", r.get("Target", ""))
                for r in relationships
                if r.get("TargetMode") != "External"
            }
            sheet_node = book.find(NS + "sheets")
            sheets = list(sheet_node) if sheet_node is not None else []
            if len(sheets) > self.settings.max_spreadsheet_sheets:
                raise DocumentLimitExceededError("Workbook sheet limit exceeded.")
            shared = []
            if "xl/sharedStrings.xml" in archive.namelist():
                with archive.open("xl/sharedStrings.xml") as stream:
                    for _, node in XML.iterparse(stream, events=("end",)):
                        if node.tag == NS + "si":
                            shared.append("".join(t.text or "" for t in node.iter(NS + "t")))
                            if len(shared) > self.settings.max_spreadsheet_cells:
                                raise DocumentLimitExceededError(
                                    "Workbook shared-string limit exceeded."
                                )
                            node.clear()
            formats: dict[int, str] = dict(BUILTIN_FORMATS)
            styles = []
            if "xl/styles.xml" in archive.namelist():
                style_root = XML.fromstring(archive.read("xl/styles.xml"))
                for node in style_root.findall(f"{NS}numFmts/{NS}numFmt"):
                    formats[int(node.get("numFmtId", "0"))] = node.get("formatCode", "")
                styles = [
                    formats.get(int(n.get("numFmtId", "0")), "")
                    for n in style_root.findall(f"{NS}cellXfs/{NS}xf")
                ]
            props = book.find(NS + "workbookPr")
            epoch = (
                "1904" if props is not None and props.get("date1904") in {"1", "true"} else "1900"
            )
            media = []
            notices = []
            # Embedded media is retained as workbook context; precise drawing anchors
            # are not fabricated. Chart-only content remains an explicit limitation.
            for entry in entries:
                if entry.filename.startswith("xl/media/"):
                    from PIL import Image

                    try:
                        data = archive.read(entry.filename)
                        with Image.open(io.BytesIO(data)) as picture:
                            if picture.width * picture.height > self.settings.max_image_pixels:
                                raise DocumentLimitExceededError(
                                    "Workbook image pixel limit exceeded."
                                )
                            out = io.BytesIO()
                            picture.convert("RGB").save(out, format="PNG")
                        media.append(
                            self.store.put_bytes(
                                f"derived/{self.document.document_id}/media_{len(media)}.png",
                                out.getvalue(),
                            )
                        )
                    except DocumentLimitExceededError:
                        raise
                    except Exception:
                        notices.append("An embedded workbook image could not be decoded.")
            if any(i.filename.startswith("xl/charts/") for i in entries):
                notices.append(
                    "Chart rendering is unavailable; chart-only information is unresolved."
                )
            if media:
                notices.append(
                    "Embedded images are workbook context; drawing anchors are unavailable."
                )
            evaluator = self._formula_index(archive, sheets, rels, shared)
            if evaluator is None:
                notices.append(
                    "Formula recalculation unavailable: workbook calculation index limit exceeded."
                )
            for index, sheet in enumerate(sheets, 1):
                part = rels.get(sheet.get(REL + "id"))
                if not part or "worksheets/" not in part:
                    self._emit(
                        index,
                        sheet.get("name", "Sheet"),
                        sheet.get("state", "visible"),
                        [],
                        ["Non-worksheet content cannot be rendered."],
                    )
                    continue

                def records(part: str = part, sheet_name: str = sheet.get("name", "Sheet")):
                    shared_formulas: dict[str, tuple[str, str]] = {}
                    with archive.open(part) as stream:
                        for _, node in XML.iterparse(stream, events=("end",)):
                            if node.tag == NS + "c":
                                ref = node.get("r", "")
                                coordinate_to_tuple(ref)
                                formula = node.find(NS + "f")
                                formula_text = self._formula_text(formula, ref, shared_formulas)
                                cached = node.find(NS + "v")
                                kind = node.get("t", "n")
                                value = cached.text if cached is not None else None
                                if kind == "s" and value is not None:
                                    value = shared[int(value)]
                                elif kind == "inlineStr":
                                    value = "".join(t.text or "" for t in node.iter(NS + "t"))
                                elif kind == "str" and cached is not None and value is None:
                                    value = ""
                                # Styles and rectangular dimensions do not create content.
                                if formula is not None or value not in (None, ""):
                                    style = int(node.get("s", "0"))
                                    fmt = styles[style] if style < len(styles) else ""
                                    cached_value = value
                                    calculation = None
                                    if formula is not None:
                                        calculation = {
                                            "status": "unsupported",
                                            "reason": "Calculation index limit exceeded.",
                                        }
                                        if evaluator is not None:
                                            try:
                                                computed, calculation = evaluator.evaluate(
                                                    sheet_name, ref
                                                )
                                                value = (
                                                    str(computed)
                                                    if isinstance(computed, Decimal)
                                                    else computed
                                                )
                                                if isinstance(value, bool):
                                                    value, kind = ("1" if value else "0"), "b"
                                                elif isinstance(computed, Decimal):
                                                    kind = "n"
                                                else:
                                                    kind = "str"
                                            except UnsupportedFormula as error:
                                                calculation = {
                                                    "status": "unsupported",
                                                    "reason": str(error),
                                                }
                                        calculation["cached_value"] = cached_value
                                    display = formatted_value(value, kind, fmt)
                                    content = (
                                        f"{ref}: {display if display is not None else '<missing>'}"
                                    )
                                    if display != value:
                                        content += f" [raw_numeric={value}]"
                                    if formula is not None:
                                        content += (
                                            f" [formula={formula_text or '<unavailable>'}; "
                                            f"cached={json.dumps(cached_value)}; "
                                            "calculation="
                                            + json.dumps(calculation, ensure_ascii=True)
                                            + "]"
                                        )
                                    if fmt and fmt != "General":
                                        content += f" [format={fmt}; date_epoch={epoch}]"
                                    yield {
                                        "cell": ref,
                                        "value": value,
                                        "display_value": display,
                                        "type": kind,
                                        "formula": formula_text,
                                        "formula_attributes": dict(formula.attrib)
                                        if formula is not None
                                        else {},
                                        "calculation": calculation,
                                        "cache_present": cached is not None,
                                        "number_format": fmt,
                                        "date_epoch": epoch,
                                        "text": content,
                                    }
                                node.clear()
                            elif node.tag == NS + "row":
                                node.clear()

                self._chunks(
                    index,
                    sheet.get("name", "Sheet"),
                    sheet.get("state", "visible"),
                    records(),
                    notices,
                    media if index == 1 else None,
                )

    @staticmethod
    def _formula_text(formula, ref, shared_formulas):
        if formula is None:
            return None
        if formula.get("t") == "shared":
            group = formula.get("si", "")
            if formula.text:
                shared_formulas[group] = (ref, formula.text)
            elif group in shared_formulas:
                origin, text = shared_formulas[group]
                try:
                    return Translator("=" + text, origin=origin).translate_formula(ref)[1:]
                except Exception:
                    return "<unsupported shared formula>"
        return formula.text or "<unsupported formula>"

    def _formula_index(self, archive, sheets, rels, shared):
        cells = {}
        for sheet in sheets:
            part = rels.get(sheet.get(REL + "id"))
            if not part or "worksheets/" not in part:
                continue
            shared_formulas = {}
            with archive.open(part) as stream:
                for _, node in XML.iterparse(stream, events=("end",)):
                    if node.tag == NS + "c":
                        ref = node.get("r", "")
                        formula = node.find(NS + "f")
                        cached = node.find(NS + "v")
                        value = cached.text if cached is not None else None
                        kind = node.get("t", "n")
                        if kind == "s" and value is not None:
                            value = shared[int(value)]
                        elif kind == "inlineStr":
                            value = "".join(t.text or "" for t in node.iter(NS + "t"))
                        elif kind == "n" and value is not None:
                            try:
                                value = Decimal(value)
                            except InvalidOperation:
                                value = "<invalid numeric value>"
                        elif kind == "b":
                            value = value in {"1", "true"}
                        elif kind == "e":
                            value = "<cached formula error>"
                        if formula is not None or value is not None:
                            if len(cells) >= self.settings.max_formula_cells:
                                return None
                            cells[(sheet.get("name", "Sheet"), ref)] = FormulaCell(
                                value,
                                self._formula_text(formula, ref, shared_formulas),
                                "Array/data-table formulas are unsupported."
                                if formula is not None
                                and formula.get("t") not in {None, "normal", "shared"}
                                else "Source cell contains an Excel error or invalid number."
                                if kind == "e" or value == "<invalid numeric value>"
                                else None,
                            )
                        node.clear()
                    elif node.tag == NS + "row":
                        node.clear()
        return FormulaEvaluator(cells, max_steps=self.settings.max_formula_steps)

    def _calamine(self, raw: bytes) -> None:
        from python_calamine import CalamineWorkbook, SheetTypeEnum

        try:
            with CalamineWorkbook.from_filelike(io.BytesIO(raw)) as book:
                if len(book.sheets_metadata) > self.settings.max_spreadsheet_sheets:
                    raise DocumentLimitExceededError("Workbook sheet limit exceeded.")
                for index, metadata in enumerate(book.sheets_metadata, 1):
                    warnings = [
                        "Binary workbook formula text, styles and drawings are unavailable; "
                        "cached values are used."
                    ]
                    if metadata.typ != SheetTypeEnum.WorkSheet:
                        self._emit(index, metadata.name, str(metadata.visible), [], warnings)
                        continue
                    sheet = book.get_sheet_by_name(metadata.name)
                    start_row, start_col = sheet.start or (0, 0)

                    def records(sheet: Any = sheet, sr: int = start_row, sc: int = start_col):
                        for row_index, row in enumerate(sheet.iter_rows(), sr + 1):
                            for col_index, value in enumerate(row, sc + 1):
                                if value is None or value == "":
                                    continue
                                if isinstance(value, (datetime, date, time)):
                                    value = value.isoformat()
                                elif isinstance(value, timedelta):
                                    value = str(value)
                                ref = f"{get_column_letter(col_index)}{row_index}"
                                yield {"cell": ref, "value": value, "text": f"{ref}: {value}"}

                    self._chunks(index, metadata.name, str(metadata.visible), records(), warnings)
        except DocumentLimitExceededError:
            raise
        except Exception as exc:
            if "password" in type(exc).__name__.lower():
                raise EncryptedDocumentError(
                    "Encrypted workbooks are unsupported; upload a decrypted copy."
                ) from None
            raise MalformedDocumentError("Binary workbook could not be decoded.") from None
