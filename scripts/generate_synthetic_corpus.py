"""Generate 20 original stress fixtures and parameter-derived draft answer keys."""

import json
import shutil
import subprocess
import time
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "eval/downloads/corpus/synthetic"
CACHE = ROOT / "eval/cache/synthetic"


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def parameters(index, kind):
    split = "development" if index < 3 else "calibration" if index == 3 else "holdout"
    names = {
        "development": "Operating expense review",
        "calibration": "Laboratory cost record",
        "holdout": "Transport settlement statement",
    }
    amount = 1200 + index * 137 + {"native": 0, "scan": 21, "docx": 42, "mixed": 63}[kind]
    return {
        "record_id": f"TEST-{kind.upper()}-{index + 1:02d}",
        "title": names[split],
        "period": "2026-Q3",
        "currency": ["EUR", "GBP", "USD", "EUR", "GBP"][index],
        "net_total": amount,
        "approved_total": amount + 75,
        "issue_date": f"2026-09-{index + 11:02d}",
        "split": split,
        "kind": kind,
        "index": index,
        "inspection_status": "PASS",
    }


def header(canvas, params, title=None):
    canvas.setFillColor(colors.HexColor("#142b46"))
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawString(54, 731, title or params["title"])
    canvas.setFillColor(colors.black)
    canvas.setFont("Helvetica", 10)
    canvas.drawString(54, 708, "Fictional evaluation document - no real person or transaction")
    canvas.drawString(54, 48, f"AI-IDP synthetic fixture | {params['record_id']}")


def lines(canvas, content, start=670):
    canvas.setFont("Helvetica", 12)
    for offset, text in enumerate(content):
        canvas.drawString(54, start - 23 * offset, text)


def summary_lines(params):
    return [
        f"Record identifier: {params['record_id']}",
        f"Issue date: {params['issue_date']}",
        f"Current reporting period: {params['period']}",
        f"Currency: {params['currency']}",
        f"Net total for current period: {params['net_total']}",
        f"Approved total for current period: {params['approved_total']}",
    ]


def draw_table(canvas, params):
    canvas.setFillColor(colors.HexColor("#e9eff5"))
    canvas.rect(54, 420, 504, 28, fill=1, stroke=0)
    canvas.setFillColor(colors.black)
    canvas.setFont("Helvetica-Bold", 11)
    for x, text in [(64, "Description"), (355, "Period"), (455, "Amount")]:
        canvas.drawString(x, 430, text)
    canvas.setFont("Helvetica", 11)
    for y, description, period, value in [
        (400, "Service allocation", params["period"], params["net_total"] - 200),
        (372, "Materials allocation", params["period"], 200),
        (344, "Historical reference only", "2025-Q3", 8900),
    ]:
        canvas.drawString(64, y, description)
        canvas.drawString(355, y, period)
        canvas.drawString(455, y, str(value))
    canvas.setStrokeColor(colors.HexColor("#d9d9d9"))
    for y in [448, 420, 390, 362, 334]:
        canvas.line(54, y, 558, y)


def native_pdf(path, params):
    canvas = Canvas(str(path), pagesize=letter, invariant=1)
    header(canvas, params)
    if params["split"] == "development":
        lines(canvas, summary_lines(params))
    elif params["split"] == "calibration":
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawString(54, 665, "Laboratory batch accounting")
        canvas.setFont("Helvetica", 11)
        for index, text in enumerate(summary_lines(params)):
            x, y = 54 + (index % 2) * 255, 625 - (index // 2) * 48
            label, value = text.split(": ", 1)
            canvas.drawString(x, y, label)
            canvas.drawString(x, y - 18, value)
    else:
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawString(
            54, 668, f"Final settlement  {params['approved_total']} {params['currency']}"
        )
        lines(canvas, summary_lines(params), start=620)
    draw_table(canvas, params)
    lines(
        canvas,
        [
            "Context",
            "The historical row is excluded from the current total.",
            "The approved total includes an additional adjustment of 75.",
            "Previous year currency: JPY. This is not the current currency.",
        ],
        start=290,
    )
    canvas.save()


def rasterize(pdf, image_path, index):
    poppler = shutil.which("pdftoppm")
    if not poppler:
        raise RuntimeError("Synthetic raster fixtures require Poppler pdftoppm for preparation")
    subprocess.run(
        [
            poppler,
            "-scale-to",
            "1800",
            "-singlefile",
            "-png",
            str(pdf),
            str(image_path.with_suffix("")),
        ],
        check=True,
        capture_output=True,
    )
    with Image.open(image_path) as original:
        image = original.convert("RGB")
        if index in [1, 3]:
            image = image.rotate(1.1, resample=Image.Resampling.BICUBIC, fillcolor="white")
        if index in [2, 4]:
            image = image.filter(ImageFilter.GaussianBlur(0.6))
        if index == 4:
            image = image.rotate(90, expand=True)
        image.save(image_path)


def scan_pdf(path, params):
    original = CACHE / (path.stem + "-native.pdf")
    image_path = CACHE / (path.stem + "-scan.png")
    native_pdf(original, params)
    rasterize(original, image_path, params["index"])
    with Image.open(image_path) as image:
        image.convert("RGB").save(
            path,
            "PDF",
            resolution=160,
            creationDate=time.gmtime(1767225600),
            modDate=time.gmtime(1767225600),
        )


def mixed_pdf(path, params):
    canvas = Canvas(str(path), pagesize=letter, invariant=1)
    header(canvas, params, "Draft estimate for comparison")
    lines(
        canvas,
        [
            f"Record identifier: {params['record_id']}",
            "DRAFT - superseded on page 3",
            "Reporting period: 2025-Q3",
            "Currency: JPY",
            "Draft approved total: 9876",
            "The final settlement must be read together with the following pages.",
        ],
    )
    canvas.showPage()
    header(canvas, params, "Allocation schedule continued")
    lines(
        canvas,
        [
            f"Record identifier: {params['record_id']}",
            f"Current reporting period: {params['period']}",
            "The current period amounts continue in the final approval.",
        ],
    )
    draw_table(canvas, params)
    canvas.showPage()
    header(canvas, params, "Final approval")
    # The answer on page 3 is deliberately raster-only; pages 1 and 2 have native text.
    image_pdf = CACHE / (path.stem + "-approval.pdf")
    image = CACHE / (path.stem + "-approval.png")
    native_pdf(image_pdf, dict(params, title="Final approval for current period"))
    rasterize(image_pdf, image, 0)
    canvas.drawImage(str(image), 0, 0, width=612, height=792)
    canvas.save()


def normalize_docx(path):
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name in sorted(members):
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, members[name])


def word_document(path, params):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.75)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.underline = False
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    document.core_properties.author = "AI-IDP synthetic fixture generator"
    document.core_properties.created = document.core_properties.modified = datetime(2026, 1, 1)
    document.add_paragraph(params["title"], "Title")
    document.add_paragraph(
        "Fictional evaluation document with a table and an embedded inspection label."
    )
    if params["split"] == "development":
        document.add_heading("Current period summary", 1)
        for text in summary_lines(params):
            document.add_paragraph(text)
    elif params["split"] == "calibration":
        document.add_heading("Laboratory batch register", 1)
        register = document.add_table(rows=0, cols=2)
        register.style = "Light Shading Accent 1"
        for text in summary_lines(params):
            label, value = text.split(": ", 1)
            cells = register.add_row().cells
            cells[0].text, cells[1].text = label, value
    else:
        document.add_heading("Settlement authorisation", 1)
        document.add_paragraph(
            f"The approved settlement for {params['period']} is {params['approved_total']} "
            f"{params['currency']}. The net total is {params['net_total']}. "
            f"Record identifier: {params['record_id']}. Issue date: {params['issue_date']}."
        )
    document.add_heading("Allocation schedule", 1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    for cell, text, width in zip(
        table.rows[0].cells, ["Description", "Period", "Amount"], [3.7, 1.4, 1.1], strict=True
    ):
        cell.text, cell.width = text, Inches(width)
    for description, value in [
        ("Service allocation", params["net_total"] - 200),
        ("Materials allocation", 200),
    ]:
        row = table.add_row()
        for cell, text in zip(row.cells, [description, params["period"], str(value)], strict=True):
            cell.text = text
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        item = OxmlElement("w:" + edge)
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            item.set(qn("w:" + key), value)
        borders.append(item)
    table._tbl.tblPr.append(borders)
    from PIL import ImageDraw, ImageFont

    label = Image.new("RGB", (700, 110), "white")
    draw = ImageDraw.Draw(label)
    font = ImageFont.load_default(size=38)
    draw.text((18, 30), "Inspection status  PASS", fill="black", font=font)
    label_path = CACHE / (path.stem + "-label.png")
    label.save(label_path)
    document.add_paragraph("Inspection label")
    document.add_picture(str(label_path), width=Inches(4.2))
    document.add_paragraph(
        "Historical comparison: 2025-Q3; total 8900 JPY. Excluded from current totals."
    )
    document.save(path)
    normalize_docx(path)


def main():
    import hashlib

    OUTPUT.mkdir(parents=True, exist_ok=True)
    CACHE.mkdir(parents=True, exist_ok=True)
    records = []
    for kind, builder, ext in [
        ("native", native_pdf, "pdf"),
        ("scan", scan_pdf, "pdf"),
        ("docx", word_document, "docx"),
        ("mixed", mixed_pdf, "pdf"),
    ]:
        for index in range(5):
            params = parameters(index, kind)
            case_id = f"synthetic-{kind}-{index + 1:02d}"
            path = OUTPUT / f"{case_id}.{ext}"
            builder(path, params)
            fixture = ROOT / "eval/fixtures/synthetic" / path.name
            fixture.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, fixture)
            answer_path = OUTPUT / f"{case_id}.parameters.json"
            write_json(answer_path, params)
            shutil.copyfile(answer_path, fixture.with_suffix(".parameters.json"))
            records.append(
                {
                    "case_id": case_id,
                    "source": "synthetic",
                    "format": ext,
                    "path": path.relative_to(ROOT).as_posix(),
                    "repository_fixture": fixture.relative_to(ROOT).as_posix(),
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "split": params["split"],
                    "group_id": f"synthetic-template:{kind}:{params['split']}",
                    "source_url": None,
                    "source_revision": "generator-v1",
                    "license": "Apache-2.0",
                    "license_url": "https://www.apache.org/licenses/LICENSE-2.0",
                    "attribution": "Original AI-IDP fictional evaluation fixture.",
                    "rights_status": "project_owned_synthetic",
                    "status": "generated",
                    "evaluation_ready": False,
                    "parameters": answer_path.relative_to(ROOT).as_posix(),
                    "tags": [kind, "missing_field", "period_distractor", "table"]
                    + (["cross_page", "mixed_content", "contradiction"] if kind == "mixed" else [])
                    + (["embedded_image"] if kind == "docx" else []),
                }
            )
            print(f"Generated {case_id}", flush=True)
    write_json(ROOT / "eval/synthetic-selection.json", records)


if __name__ == "__main__":
    main()
