"""Project-owned workbook stress fixture; generated large files stay outside git."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill


def workbook_bytes(rows: int = 80, sheets: int = 2) -> bytes:
    import io

    book = Workbook()
    book.remove(book.active)
    for sheet_index in range(sheets):
        sheet = book.create_sheet(f"Register {sheet_index + 1}")
        sheet.append(["Record ID", "Description", "Quantity", "Amount EUR"])
        for index in range(sheet_index * rows // sheets, (sheet_index + 1) * rows // sheets):
            sheet.append(
                [
                    f"LOT-{index:06d}",
                    f"Synthetic item {index:06d}",
                    index % 9 + 1,
                    (index + 1) * 1.25,
                ]
            )
        # A distant formatted blank must not turn into a million model-input rows.
        sheet["Z1000000"].fill = PatternFill("solid", fgColor="EEEEEE")
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 32
        sheet.column_dimensions["D"].width = 18
    book.properties.creator = "AI-IDP"
    book.properties.lastModifiedBy = "AI-IDP"
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, default=80)
    parser.add_argument("--sheets", type=int, default=2)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("eval/downloads/spreadsheet-research/synthetic-register.xlsx"),
    )
    args = parser.parse_args()
    if args.rows < 1 or not 1 <= args.sheets <= args.rows:
        parser.error("Require rows >= sheets >= 1.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(workbook_bytes(args.rows, args.sheets))
    print(f"Generated {args.rows} synthetic records across {args.sheets} sheets.")


if __name__ == "__main__":
    main()
