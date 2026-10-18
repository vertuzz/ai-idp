"""Opt-in spreadsheet parser research; never calls models or modifies the frozen corpus.

Optional packages are supplied with ``uv run --locked --with ...``; see the research plan.
Pinned source hashes are tracked; downloaded bytes and measurements remain ignored.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import importlib.metadata
import json
import os
import platform
import subprocess
import sys
import threading
import time
import warnings
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "eval/downloads/spreadsheet-research"
RESULTS = ROOT / "eval/cache/spreadsheet-research"
CATALOG = ROOT / "eval/spreadsheet-research-sources.json"


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def acquire() -> None:
    import httpx

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    with httpx.Client(follow_redirects=True, timeout=60) as client:
        for item in json.loads(CATALOG.read_text(encoding="utf-8"))["sources"]:
            name = item["filename"]
            path = DOWNLOADS / name
            if not path.exists():
                response = client.get(item["url"])
                response.raise_for_status()
                if hashlib.sha256(response.content).hexdigest() != item["sha256"]:
                    raise ValueError(f"Upstream source changed: {name}; no replacement saved.")
                path.write_bytes(response.content)
    verified_sources()


def verified_sources() -> list[Path]:
    paths = []
    for item in json.loads(CATALOG.read_text(encoding="utf-8"))["sources"]:
        path = DOWNLOADS / item["filename"]
        if not path.exists():
            raise ValueError(f"Missing source: {item['filename']}; use --download explicitly.")
        if hashlib.sha256(path.read_bytes()).hexdigest() != item["sha256"]:
            raise ValueError(f"Source changed: {item['filename']}; no replacement performed.")
        paths.append(path)
    return sorted(paths)


def parse_file(path: Path, engine: str) -> dict[str, Any]:
    import psutil

    process = psutil.Process()
    peak = process.memory_info().rss
    done = threading.Event()

    def sample() -> None:
        nonlocal peak
        while not done.wait(0.01):
            peak = max(peak, process.memory_info().rss)

    sampler = threading.Thread(target=sample, daemon=True)
    sampler.start()
    started = time.perf_counter()
    result: dict[str, Any] = {"filename": path.name, "engine": engine}
    try:
        count = 0
        sheets = []
        with warnings.catch_warnings(record=True) as notices:
            if engine == "openpyxl":
                import openpyxl

                book = openpyxl.load_workbook(
                    path,
                    read_only=True,
                    data_only=False,
                    keep_links=False,
                )
                for sheet in book:
                    populated = sum(
                        cell.value is not None for row in sheet.iter_rows() for cell in row
                    )
                    count += populated
                    sheets.append({"name": sheet.title, "nonempty_cells": populated})
                book.close()
            else:
                from python_calamine import CalamineWorkbook, SheetTypeEnum

                with CalamineWorkbook.from_path(str(path)) as book:
                    for metadata in book.sheets_metadata:
                        name = metadata.name
                        if metadata.typ != SheetTypeEnum.WorkSheet:
                            sheets.append({"name": name, "unsupported_kind": str(metadata.typ)})
                            continue
                        sheet = book.get_sheet_by_name(name)
                        populated = sum(
                            value not in (None, "") for row in sheet.iter_rows() for value in row
                        )
                        count += populated
                        sheets.append({"name": name, "nonempty_cells": populated})
            result.update(
                status="parsed",
                nonempty_cells=count,
                sheets=sheets,
                warnings=sorted({str(w.message) for w in notices}),
            )
    except Exception as exc:
        result.update(status="error", error_type=type(exc).__name__)
    finally:
        done.set()
        sampler.join()
        peak = max(peak, process.memory_info().rss)
    result.update(
        seconds=round(time.perf_counter() - started, 3), peak_rss_mib=round(peak / 2**20, 1)
    )
    return result


def inspect_ooxml(path: Path) -> dict[str, Any]:
    from xml.etree import ElementTree as ET

    if not zipfile.is_zipfile(path):
        return {"filename": path.name, "status": "not-zip; encryption or invalid container"}
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = book.find("s:sheets", ns)
        formulas = empty_string_cache = missing_cache = cells = 0
        for name in names:
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                with archive.open(name) as stream:
                    for _, node in ET.iterparse(stream, events=("end",)):
                        if node.tag == "{" + ns["s"] + "}c":
                            cells += 1
                            if node.find("s:f", ns) is not None:
                                formulas += 1
                                cached = node.find("s:v", ns)
                                if cached is None:
                                    missing_cache += 1
                                elif cached.text is None:
                                    if node.get("t") == "str":
                                        empty_string_cache += 1
                                    else:
                                        missing_cache += 1
                            node.clear()
        return {
            "filename": path.name,
            "bytes": path.stat().st_size,
            "expanded_bytes": sum(x.file_size for x in archive.infolist()),
            "xml_cells_including_styled_blanks": cells,
            "formulas": formulas,
            "formulas_with_empty_string_cache": empty_string_cache,
            "formulas_without_cached_value": missing_cache,
            "vba_parts": sum(n.lower().endswith("vbaproject.bin") for n in names),
            "image_parts": sum(n.startswith("xl/media/") for n in names),
            "sheets": [dict(s.attrib) for s in sheets] if sheets is not None else [],
        }


async def sdk_probe() -> list[dict[str, Any]]:
    os.environ.setdefault("PYDANTIC_AI_NO_BANNER", "1")
    from pydantic_ai import Agent, StructuredDict
    from pydantic_ai.messages import ModelResponse, ToolCallPart
    from pydantic_ai.models.function import FunctionModel

    schema = {
        "type": "object",
        "properties": {"items": {"type": "array", "items": {"type": "integer"}}},
        "required": ["items"],
        "additionalProperties": False,
    }
    records = []
    for mode in ("truncated_json", "valid_json_length", "valid_json_stop"):
        calls = 0

        def respond(messages: Any, info: Any, mode: str = mode) -> ModelResponse:
            nonlocal calls
            calls += 1
            args = '{"items":[1,2' if mode == "truncated_json" else '{"items":[1,2]}'
            return ModelResponse(
                parts=[ToolCallPart(info.output_tools[0].name, args)],
                finish_reason="stop" if mode == "valid_json_stop" else "length",
            )

        agent = Agent(
            FunctionModel(respond),
            output_type=StructuredDict(schema),
            retries=3,
            model_settings={"max_tokens": 8192},
        )
        try:
            output = await agent.run("Return all items.")
            records.append(
                {"mode": mode, "calls": calls, "status": "accepted", "output": output.output}
            )
        except Exception as exc:
            records.append({"mode": mode, "calls": calls, "status": type(exc).__name__})
    return records


def output_probe() -> dict[str, Any]:
    import tiktoken

    from ai_idp.domain.field_paths import concrete_targets

    schema = {
        "type": "object",
        "properties": {
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "sku": {"type": "string"},
                        "description": {"type": "string"},
                        "quantity": {"type": "integer"},
                        "amount": {"type": "number"},
                    },
                    "required": ["sku", "description", "quantity", "amount"],
                },
            }
        },
    }
    records = []
    enc = tiktoken.get_encoding("cl100k_base")
    for size in (25, 50, 100, 250, 399, 400, 1000, 10000):
        items = [
            {
                "sku": f"SKU-{i:06d}",
                "description": f"Synthetic item {i:06d}",
                "quantity": i % 9 + 1,
                "amount": 125.25,
            }
            for i in range(size)
        ]
        evidence = [
            {
                "field_path": f"/items/{i}/{key}",
                "unit_id": "sheet_0001_range_0001",
                "quote": str(value),
                "raw_value": str(value),
                "text_span": None,
            }
            for i, item in enumerate(items)
            for key, value in item.items()
        ]
        data = {"items": items}
        envelope = {**data, "_evidence": evidence}
        try:
            targets = len(concrete_targets(schema, data))
        except Exception as exc:
            targets = type(exc).__name__
        compact = json.dumps(envelope, ensure_ascii=False, separators=(",", ":"))
        records.append(
            {
                "items": size,
                "targets": targets,
                "data_proxy_tokens": len(enc.encode(json.dumps(data))),
                "envelope_proxy_tokens": len(enc.encode(compact)),
                "envelope_bytes": len(compact.encode()),
            }
        )
    return {
        "tokenizer": "cl100k_base; size proxy, not provider-native billing tokens",
        "rows": records,
        "sdk": asyncio.run(sdk_probe()),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download", action="store_true")
    parser.add_argument("--child", nargs=2, metavar=("ENGINE", "FILE"))
    args = parser.parse_args()
    if args.child:
        print(json.dumps(parse_file(Path(args.child[1]), args.child[0])))
        return
    if args.download:
        acquire()
    paths = verified_sources()
    records = []
    for path in paths:
        for engine in ("openpyxl", "calamine"):
            if engine == "openpyxl" and path.suffix not in {".xlsx", ".xlsm"}:
                continue
            try:
                child = subprocess.run(
                    [
                        sys.executable,
                        "-m",
                        "scripts.research_spreadsheets",
                        "--child",
                        engine,
                        str(path),
                    ],
                    cwd=ROOT,
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=True,
                )
                result = json.loads(child.stdout)
            except (subprocess.TimeoutExpired, subprocess.CalledProcessError, ValueError) as exc:
                result = {"filename": path.name, "engine": engine, "status": type(exc).__name__}
            records.append(result)
            print(
                json.dumps({k: v for k, v in result.items() if k not in {"sheets", "warnings"}}),
                flush=True,
            )
    report = {
        "python": platform.python_version(),
        "platform": platform.system(),
        "versions": {
            p: importlib.metadata.version(p)
            for p in ("openpyxl", "python-calamine", "tiktoken", "pydantic-ai-slim")
        },
        "parsers": records,
        "ooxml": [inspect_ooxml(p) for p in paths if p.suffix in {".xlsx", ".xlsm"}],
        "output": output_probe(),
        "limitations": (
            "Research only; no live model calls, quality benchmark or formula calculation."
        ),
    }
    write_json(RESULTS / "report.json", report)
    print(json.dumps(report["output"], indent=2))


if __name__ == "__main__":
    main()
