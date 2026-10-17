"""Validate corpus integrity and report annotation readiness separately from file coverage."""

import argparse
import json
from collections import Counter, defaultdict

from docx import Document
from PIL import Image
from pypdf import PdfReader

from scripts.prepare_cord_preview import ROOT, sha256, write_json

QUOTAS = {"doclaynet": 35, "cord_v2": 30, "gov_uk_pdf": 20, "gov_uk_docx": 15, "synthetic": 20}
SPLITS = {"development": 72, "calibration": 24, "holdout": 24}


def safe_path(root, relative):
    path = (root / relative).resolve()
    if root.resolve() not in path.parents:
        raise ValueError(f"Corpus path is outside the repository: {relative}")
    return path


def validate(cases, root=ROOT):
    errors = []
    groups, checksums = defaultdict(set), defaultdict(list)
    units, sizes, formats = Counter(), Counter(), Counter()
    if Counter(c["source"] for c in cases) != QUOTAS:
        errors.append("Source quotas do not match 35/30/20/15/20")
    if Counter(c["split"] for c in cases) != SPLITS:
        errors.append("Split counts do not match 72/24/24")
    if len({c["case_id"] for c in cases}) != len(cases):
        errors.append("Case IDs are not unique")
    for case in cases:
        name = case["case_id"]
        try:
            path = safe_path(root, case["path"])
            if sha256(path) != case["sha256"]:
                raise ValueError("SHA-256 mismatch")
            if path.stat().st_size != case["bytes"]:
                raise ValueError("Byte count mismatch")
            if not case.get("attribution") or not case.get("license_url"):
                raise ValueError("Missing licence/attribution metadata")
            for asset in case.get("assets", {}).values():
                if sha256(safe_path(root, asset["path"])) != asset["sha256"]:
                    raise ValueError("Derived/source counterpart checksum mismatch")
            if "upstream_annotation_sha256" in case:
                if (
                    sha256(safe_path(root, case["upstream_annotation"]))
                    != case["upstream_annotation_sha256"]
                ):
                    raise ValueError("Upstream annotation checksum mismatch")
            groups[case["group_id"]].add(case["split"])
            checksums[case["sha256"]].append(name)
            formats[case["format"]] += 1
            sizes[case["source"]] += path.stat().st_size
            if case["format"] == "pdf":
                reader = PdfReader(path)
                pages = len(reader.pages)
                if pages == 0:
                    raise ValueError("PDF has no pages")
                if case["source"] == "doclaynet" and pages != 1:
                    raise ValueError("DocLayNet original is not a single-page PDF")
                for page in reader.pages:
                    if float(page.mediabox.width) <= 0 or float(page.mediabox.height) <= 0:
                        raise ValueError("PDF page dimensions are invalid")
                units[case["source"]] += pages
            elif case["format"] == "docx":
                document = Document(path)
                if not document.paragraphs and not document.tables:
                    raise ValueError("DOCX has no readable body")
                units[case["source"]] += len(document.element.body)
            else:
                with Image.open(path) as image:
                    image.verify()
                units[case["source"]] += 1
            annotation = json.loads(
                safe_path(root, case["annotation_path"]).read_text(encoding="utf-8")
            )
            if annotation["case_id"] != name:
                raise ValueError("Annotation case ID mismatch")
            for field in annotation["expected_fields"]:
                if not field["path"].startswith("/"):
                    raise ValueError("Expected JSON Pointer field path")
            if case["evaluation_ready"] and annotation["review_status"] != "reviewed":
                raise ValueError("Unreviewed annotation was marked evaluation-ready")
        except Exception as error:
            errors.append(f"{name}: {error}")
    for group, splits in groups.items():
        if len(splits) > 1:
            errors.append(f"Source/template group crosses splits: {group}")
    for digest, ids in checksums.items():
        if len(ids) > 1:
            errors.append(f"Duplicate source bytes ({digest}): {ids}")
    return {
        "file_integrity_passed": not errors,
        "errors": errors,
        "documents": len(cases),
        "sources": dict(Counter(c["source"] for c in cases)),
        "splits": dict(Counter(c["split"] for c in cases)),
        "formats": dict(formats),
        "primary_bytes_by_source": dict(sizes),
        "source_units_by_source": dict(units),
        "unit_note": (
            "PDF pages, image frames, and DOCX top-level XML body elements; "
            "DOCX elements are not physical pages or finalized ingestion blocks."
        ),
        "annotation_status": dict(Counter(c["annotation_status"] for c in cases)),
        "evaluation_ready": sum(c["evaluation_ready"] for c in cases),
        "remaining_reviews": [
            "Human field/evidence annotation review",
            "File-specific rights and privacy review",
            "Visual/template near-duplicate adjudication beyond recorded groups",
        ],
        "selection_status": "candidate_v1_files_complete; final_benchmark_freeze_pending_reviews",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--require-ready", action="store_true")
    args = parser.parse_args()
    cases = [
        json.loads(line)
        for line in (ROOT / "eval/manifest.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    report = validate(cases)
    write_json(ROOT / "eval/corpus-status.json", report)
    print(json.dumps(report, indent=2))
    if report["errors"] or (args.require_ready and report["evaluation_ready"] != 120):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
