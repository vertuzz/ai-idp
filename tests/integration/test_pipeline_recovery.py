import asyncio
import io
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import httpx
import pytest
from docx import Document as WordDocument
from PIL import Image
from sqlalchemy import update

from ai_idp.adapters.inbound.http.app import create_app
from ai_idp.adapters.inbound.http.dependencies import build_container
from ai_idp.adapters.outbound.persistence.models import (
    DocumentUnitRow,
    FieldAttemptRow,
    JobQueueRow,
)
from ai_idp.application.ports.protocols import ExtractionBatchResult
from ai_idp.domain.enums import JobStatus, VerificationStatus
from ai_idp.domain.errors import BudgetExhaustedError, JobLeaseLostError
from ai_idp.domain.models import Evidence, UsageRecord
from ai_idp.settings import Settings
from ai_idp.workers.runner import WorkerRunner
from tests.integration.test_api_and_worker import FakeDecisionEngine, FakeExtractionEngine

ROOT = Path(__file__).resolve().parents[2]
SCHEMA = {
    "type": "object",
    "properties": {"record_id": {"type": "string"}},
    "required": ["record_id"],
}


def setup(tmp_path, *, extractor=None, decision=None, nullable=False, **limits):
    c = build_container(
        Settings(
            database_url=limits.pop(
                "database_url", f"sqlite:///{(tmp_path / 'db.sqlite3').as_posix()}"
            ),
            local_storage_root=tmp_path / "storage",
            local_telemetry_root=tmp_path / "telemetry",
            **limits,
        ),
        decision_engine=decision or FakeDecisionEngine(),
        extraction_engine=extractor or FakeExtractionEngine(),
        create_tables=True,
    )
    path = ROOT / "eval/fixtures/synthetic/synthetic-native-01.pdf"
    doc = c.ingest_use_case.execute(raw_bytes=path.read_bytes(), filename=path.name)
    schema = (
        SCHEMA
        if not nullable
        else {
            "type": "object",
            "properties": {"record_id": {"type": ["string", "null"]}},
            "required": ["record_id"],
        }
    )
    job = c.request_extraction_use_case.execute(document_id=doc.document_id, schema_json=schema)
    return c, doc, job


class ControlledExtractor:
    def __init__(self, value="TEST-NATIVE-01", cost=0.0, evidence=True):
        self.value = value
        self.cost = cost
        self.evidence = evidence
        self.calls = 0

    async def extract_fields(self, **kw):
        self.calls += 1
        return ExtractionBatchResult(
            data={f.name: self.value for f in kw["target_fields"]},
            evidence_by_field={
                f.path: [Evidence(unit_id=kw["unit_inputs"][0][0].unit.unit_id)]
                for f in kw["target_fields"]
            }
            if self.evidence
            else {},
            model_id=kw["model_id"],
            usage=UsageRecord(
                usage_id=uuid.uuid4().hex,
                job_id=kw["job_id"],
                stage=kw["stage_name"],
                model_id=kw["model_id"],
                estimated_cost_usd=self.cost,
            ),
        )


def make_available(c, job_id):
    with c.job_repo._session_factory() as db:
        db.execute(
            update(JobQueueRow)
            .where(JobQueueRow.job_id == job_id)
            .values(available_at=datetime.now(UTC) - timedelta(seconds=1))
        )
        db.commit()


async def test_crash_resume_reuses_paid_steps_and_preserves_attempt(tmp_path):
    ext = ControlledExtractor(cost=0.06)
    c, _, job = setup(tmp_path, extractor=ext, max_job_estimated_cost_usd=0.10)
    with patch.object(
        c.result_repo, "save_final_result", side_effect=RuntimeError("injected crash")
    ):
        with pytest.raises(RuntimeError, match="injected crash"):
            await c.pipeline_service.execute_job(job.job_id)
    first = c.job_repo.get_job(job.job_id)
    assert first.status == JobStatus.PENDING
    with c.result_repo._session_factory() as db:
        attempt = db.query(FieldAttemptRow).filter_by(job_id=job.job_id).one()
        original_time = attempt.created_at
        original_candidate = attempt.candidate_json
    make_available(c, job.job_id)
    result = await c.pipeline_service.execute_job(job.job_id)
    assert result.status == JobStatus.COMPLETED
    assert ext.calls == 1
    assert 0.06 <= result.usage.estimated_cost_usd < 0.10
    assert result.fields["/record_id"].attempt_count == 1
    assert result.fields["/record_id"].attempts[0].created_at.replace(
        tzinfo=None
    ) == original_time.replace(tzinfo=None)
    with c.result_repo._session_factory() as db:
        assert (
            db.query(FieldAttemptRow).filter_by(job_id=job.job_id).one().candidate_json
            == original_candidate
        )


async def test_spend_exhaustion_is_durable_and_not_requeued(tmp_path):
    ext = ControlledExtractor(cost=0.11)
    c, _, job = setup(tmp_path, extractor=ext, max_job_estimated_cost_usd=0.10)
    with pytest.raises(BudgetExhaustedError):
        await c.pipeline_service.execute_job(job.job_id)
    assert c.job_repo.get_job(job.job_id).status == JobStatus.FAILED
    assert c.job_repo.get_job(job.job_id).usage.estimated_cost_usd >= 0.11
    assert c.job_repo.lease_next_job("other", 120) is None
    assert ext.calls == 1


@pytest.mark.parametrize(
    "verdict", [VerificationStatus.CONTRADICTED, VerificationStatus.INSUFFICIENT_EVIDENCE]
)
async def test_nullable_missing_value_needs_a_real_absence_verdict(tmp_path, verdict):
    decision = FakeDecisionEngine()
    original = decision.verify_fields

    async def verify(**kwargs):
        result = await original(**kwargs)
        for outcome in result.outcomes.values():
            outcome.status = verdict
            outcome.raw_probabilities = {"SUPPORTED": 0.01, verdict.value: 0.99}
        return result

    decision.verify_fields = verify
    c, _, job = setup(
        tmp_path,
        extractor=ControlledExtractor(None),
        decision=decision,
        nullable=True,
        max_field_extra_attempts=0,
        max_coverage_reviews=0,
    )
    result = await c.pipeline_service.execute_job(job.job_id)
    assert sum(decision._verify_counts.values()) == 1
    assert result.fields["/record_id"].status == "unresolved"
    assert result.status == JobStatus.COMPLETED_WITH_UNRESOLVED


async def test_confirmed_nullable_absence_can_be_verified(tmp_path):
    c, _, job = setup(tmp_path, extractor=ControlledExtractor(None), nullable=True)
    result = await c.pipeline_service.execute_job(job.job_id)
    assert result.status == JobStatus.COMPLETED
    assert result.fields["/record_id"].verification == VerificationStatus.SUPPORTED


async def test_missing_field_in_selected_unit_gets_bounded_recovery(tmp_path):
    ext = ControlledExtractor(None)
    decision = FakeDecisionEngine(contradict_first_attempt_for={"/record_id"})
    c, _, job = setup(
        tmp_path, extractor=ext, decision=decision, nullable=True, max_field_extra_attempts=0
    )
    result = await c.pipeline_service.execute_job(job.job_id)
    assert ext.calls == 2
    assert result.fields["/record_id"].attempt_count == 2
    assert result.fields["/record_id"].verification == VerificationStatus.SUPPORTED


async def test_missing_citation_never_becomes_verified(tmp_path):
    c, _, job = setup(
        tmp_path, extractor=ControlledExtractor(evidence=False), max_field_extra_attempts=0
    )
    result = await c.pipeline_service.execute_job(job.job_id)
    field = result.fields["/record_id"]
    assert field.sources == []
    assert field.verification == VerificationStatus.INSUFFICIENT_EVIDENCE
    assert field.status == "unresolved"


def test_bad_unit_ids_are_dropped_and_canonical_locators_cannot_be_overwritten(tmp_path):
    c, doc, _ = setup(tmp_path)
    units = c.job_repo.get_units(doc.document_id)
    good = units[0].unit_id
    result = c.pipeline_service._normalize_evidences(
        [
            Evidence(unit_id="invented", quote="made up"),
            Evidence(unit_id=good, locator={"page_number": 999}, quote="original"),
        ],
        unit_by_id={u.unit_id: u for u in units},
        default_unit_ids=[good],
    )
    assert len(result) == 1 and result[0].unit_id == good
    assert result[0].locator["page_number"] == 1


async def test_supported_threshold_is_applied(tmp_path):
    c, _, job = setup(tmp_path, verification_supported_threshold=0.99, max_field_extra_attempts=0)
    result = await c.pipeline_service.execute_job(job.job_id)
    assert result.fields["/record_id"].verification == VerificationStatus.INSUFFICIENT_EVIDENCE


async def test_inline_execution_cannot_be_leased_or_executed_twice(tmp_path):
    entered = asyncio.Event()
    release = asyncio.Event()
    ext = ControlledExtractor()
    original = ext.extract_fields

    async def blocked(**kw):
        entered.set()
        await release.wait()
        return await original(**kw)

    ext.extract_fields = blocked
    c, doc, queued = setup(tmp_path, extractor=ext)
    c.job_repo.request_cancel(queued.job_id)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app(c)), base_url="http://test"
    ) as client:
        task = asyncio.create_task(
            client.post(
                "/v1/extractions?run_inline=true",
                json={"document_id": doc.document_id, "schema": SCHEMA},
            )
        )
        await asyncio.wait_for(entered.wait(), 3)
        assert await WorkerRunner(c).process_one_if_available() is None
        with c.job_repo._session_factory() as db:
            active = db.query(JobQueueRow).filter_by(status="leased").one()
            active_id = active.job_id
        with pytest.raises(JobLeaseLostError):
            await c.pipeline_service.execute_job(active_id)
        release.set()
        assert (await task).status_code == 202
    assert ext.calls == 1


def test_stale_owner_cannot_write_after_reclaim(tmp_path):
    c, _, job = setup(tmp_path)
    assert c.job_repo.claim_job(job.job_id, "old", 120)
    with c.job_repo._session_factory() as db:
        db.execute(
            update(JobQueueRow)
            .where(JobQueueRow.job_id == job.job_id)
            .values(lease_until=datetime.now(UTC) - timedelta(seconds=1))
        )
        db.commit()
    assert c.job_repo.claim_job(job.job_id, "new", 120)
    with c.job_repo.execution_scope(job.job_id, "old"):
        assert not c.job_repo.heartbeat_job(job.job_id, "old", 120)
        for write in (
            lambda: c.job_repo.complete_queue_item(job.job_id, JobStatus.COMPLETED),
            lambda: c.job_repo.update_job_progress(job.job_id, status=JobStatus.COMPLETED),
            lambda: c.result_repo.save_final_result(job.job_id, JobStatus.COMPLETED, {}),
            lambda: c.result_repo.save_step(job.job_id, "invalid", {}, None),
        ):
            with pytest.raises(JobLeaseLostError):
                write()
    with c.job_repo._session_factory() as db:
        assert db.get(JobQueueRow, job.job_id).lease_owner == "new"


async def test_worker_runs_two_jobs_concurrently(tmp_path):
    both = asyncio.Event()
    release = asyncio.Event()
    active = 0
    ext = ControlledExtractor()
    original = ext.extract_fields

    async def blocked(**kw):
        nonlocal active
        active += 1
        if active == 2:
            both.set()
        await release.wait()
        active -= 1
        return await original(**kw)

    ext.extract_fields = blocked
    c, doc, first = setup(tmp_path, extractor=ext, max_concurrent_jobs=2)
    second = c.request_extraction_use_case.execute(document_id=doc.document_id, schema_json=SCHEMA)
    worker = WorkerRunner(c)
    task = asyncio.create_task(worker.run_forever())
    try:
        await asyncio.wait_for(both.wait(), 3)
    finally:
        worker.stop()
        release.set()
        await asyncio.wait_for(task, 3)
    assert c.job_repo.get_job(first.job_id).status == JobStatus.COMPLETED
    assert c.job_repo.get_job(second.job_id).status == JobStatus.COMPLETED


async def test_lost_lease_cancels_active_inference(tmp_path):
    entered = asyncio.Event()
    cancelled = asyncio.Event()
    ext = ControlledExtractor()

    async def blocked(**kwargs):
        entered.set()
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.set()

    ext.extract_fields = blocked
    c, _, job = setup(tmp_path, extractor=ext, worker_lease_seconds=1)
    task = asyncio.create_task(c.pipeline_service.execute_job(job.job_id))
    await asyncio.wait_for(entered.wait(), 3)
    with c.job_repo._session_factory() as db:
        db.execute(
            update(JobQueueRow)
            .where(JobQueueRow.job_id == job.job_id)
            .values(
                lease_owner="replacement", lease_until=datetime.now(UTC) + timedelta(seconds=120)
            )
        )
        db.commit()
    with pytest.raises(JobLeaseLostError):
        await asyncio.wait_for(task, 3)
    assert cancelled.is_set()
    with c.job_repo._session_factory() as db:
        assert db.get(JobQueueRow, job.job_id).lease_owner == "replacement"


async def test_cancelled_execution_cannot_finalize(tmp_path):
    entered = asyncio.Event()
    release = asyncio.Event()
    ext = ControlledExtractor()
    original = ext.extract_fields

    async def blocked(**kwargs):
        entered.set()
        await release.wait()
        return await original(**kwargs)

    ext.extract_fields = blocked
    c, _, job = setup(tmp_path, extractor=ext)
    task = asyncio.create_task(c.pipeline_service.execute_job(job.job_id))
    await asyncio.wait_for(entered.wait(), 3)
    c.job_repo.request_cancel(job.job_id)
    release.set()
    result = await task
    assert result.status == JobStatus.CANCELLED
    assert c.job_repo.get_job(job.job_id).status == JobStatus.CANCELLED


async def test_docx_keeps_all_images_and_baseline_keeps_text(tmp_path):
    c, _, _ = setup(tmp_path)
    word = WordDocument()
    word.add_heading("Labels", level=1)
    for color in ["red", "blue"]:
        image = io.BytesIO()
        Image.new("RGB", (40, 40), color).save(image, format="PNG")
        image.seek(0)
        word.add_picture(image)
    buf = io.BytesIO()
    word.save(buf)
    doc = c.ingest_use_case.execute(raw_bytes=buf.getvalue(), filename="two.docx")
    units = c.job_repo.get_units(doc.document_id)
    assert len(units) == 1 and len(units[0].image_refs) == 2
    with c.job_repo._session_factory() as db:
        row = db.query(DocumentUnitRow).filter_by(document_id=doc.document_id).one()
        row.image_refs_json = []
        row.locator_json = {**row.locator_json, "embedded_image": "image1.png"}
        db.commit()
    units = c.job_repo.get_units(doc.document_id)
    assert len(units[0].image_refs) == 2
    assert "embedded_image" not in units[0].source_locator
    payload = c.pipeline_service._load_unit_payloads(units)[0]
    assert payload.image_bytes and len(payload.additional_images) == 1
    job = c.request_extraction_use_case.execute(
        document_id=doc.document_id, schema_json=SCHEMA, model_profile="always_vision"
    )
    spec = c.job_repo.get_spec(job.spec_id)
    modality = await c.pipeline_service._determine_unit_modality(
        job=job, spec=spec, payload=payload, record_usage=lambda _: None
    )
    assert modality == "TEXT_AND_VISION"
