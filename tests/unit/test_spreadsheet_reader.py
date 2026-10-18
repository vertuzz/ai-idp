import io
import zipfile

import pytest
from openpyxl import Workbook

from ai_idp.adapters.outbound.documents.spreadsheet_reader import SpreadsheetDocumentReader
from ai_idp.adapters.outbound.storage.local import LocalDocumentBlobStore
from ai_idp.application.use_cases.ingest_document import detect_media_type
from ai_idp.domain.errors import DocumentLimitExceededError
from ai_idp.domain.media_types import XLSM_MIME, XLSX_MIME
from ai_idp.domain.models import Document
from ai_idp.settings import Settings
from scripts.generate_spreadsheet_fixture import workbook_bytes


def read(tmp_path, data, **options):
    settings = Settings(_env_file=None, **options)
    store = LocalDocumentBlobStore(tmp_path)
    document = Document(
        document_id="synthetic",
        filename="test.xlsx",
        sha256="test",
        bytes_size=len(data),
        storage_key="original",
        media_type=detect_media_type(data),
    )
    return SpreadsheetDocumentReader(settings).normalize(document, data, store), store


def test_sparse_workbook_retains_records_across_sheets_without_styled_blanks(tmp_path):
    data = workbook_bytes(80, 2)
    units, store = read(tmp_path, data)
    assert len(units) == 4
    text = "\n".join(store.get_text(unit.native_text_ref) for unit in units)
    assert text.count("LOT-") == 80
    assert "Z1000000" not in text
    assert {u.source_locator["sheet_name"] for u in units} == {"Register 1", "Register 2"}
    assert detect_media_type(data) == XLSX_MIME


def test_formulas_hidden_sheets_zero_and_identifiers_remain_distinct(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(["000123", 0, "=1+1", False])
    book.create_sheet("Hidden").sheet_state = "veryHidden"
    out = io.BytesIO()
    book.save(out)
    units, store = read(tmp_path, out.getvalue())
    text = store.get_text(units[0].native_text_ref)
    assert "000123" in text and "B1: 0" in text and "formula=1+1" in text
    assert "D1: FALSE" in text
    assert units[1].metadata["sheet_state"] == "veryHidden"


def test_macro_container_is_read_without_executing_vba(tmp_path):
    original = workbook_bytes(3, 1)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(out, "w") as target:
        for name in source.namelist():
            target.writestr(name, source.read(name))
        # An inert marker tests package admission. Real VBA is separately exercised
        # using the pinned external macro workbook, never executed or committed.
        target.writestr("xl/vbaProject.bin", b"INERT TEST MARKER")
    assert detect_media_type(out.getvalue()) == XLSM_MIME
    units, store = read(tmp_path, out.getvalue())
    assert "LOT-000002" in store.get_text(units[0].native_text_ref)


def test_cell_and_sheet_limits_fail_explicitly(tmp_path):
    data = workbook_bytes(20, 2)
    with pytest.raises(DocumentLimitExceededError, match="cell"):
        read(tmp_path, data, max_spreadsheet_cells=5)
    with pytest.raises(DocumentLimitExceededError, match="sheet"):
        read(tmp_path, data, max_spreadsheet_sheets=1)


def test_simple_numeric_id_masks_and_percentages_retain_raw_values(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 123
    sheet["A1"].number_format = "000000"
    sheet["B1"] = 0.25
    sheet["B1"].number_format = "0.0%"
    out = io.BytesIO()
    book.save(out)
    units, store = read(tmp_path, out.getvalue())
    text = store.get_text(units[0].native_text_ref)
    assert "A1: 000123 [raw_numeric=123]" in text
    assert "B1: 25.0% [raw_numeric=0.25]" in text
