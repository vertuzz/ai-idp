import io
import zipfile

from docx import Document as WordDocument
from PIL import Image

from ai_idp.adapters.outbound.documents.docx_reader import MammothDocxDocumentReader
from ai_idp.adapters.outbound.storage.local import LocalDocumentBlobStore
from ai_idp.application.ports.protocols import UnitInputPayload, VerificationBatchResult
from ai_idp.application.services.pipeline import ExtractionPipelineService
from ai_idp.domain.enums import UnitKind, VerificationStatus
from ai_idp.domain.models import Document, DocumentUnit, VerificationOutcome
from ai_idp.domain.schema_normalizer import normalize_extraction_schema


def test_unsupported_embedded_vector_is_preserved_and_never_mislabeled_png(tmp_path):
    picture = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(picture, format="PNG")
    picture.seek(0)
    document = WordDocument()
    document.add_paragraph("Useful semantic text remains available.")
    document.add_picture(picture)
    source = io.BytesIO()
    document.save(source)
    changed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source.getvalue())) as original:
        with zipfile.ZipFile(changed, "w") as archive:
            for entry in original.infolist():
                content = original.read(entry.filename)
                if entry.filename == "[Content_Types].xml":
                    content = content.replace(
                        b'ContentType="image/png"', b'ContentType="image/x-emf"'
                    )
                if entry.filename.startswith("word/media/"):
                    content = b"EMF_SOURCE_BYTES"
                archive.writestr(entry, content)
    store = LocalDocumentBlobStore(tmp_path)
    units = MammothDocxDocumentReader().normalize(
        Document(
            document_id="vector",
            filename="v.docx",
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            sha256="unused",
            bytes_size=len(changed.getvalue()),
            storage_key="source",
        ),
        changed.getvalue(),
        store,
    )
    assert any(u.native_text_ref for u in units)
    unavailable = [ref for u in units for ref in u.metadata["unavailable_images"]]
    assert unavailable
    assert store.get_bytes(unavailable[0]) == b"EMF_SOURCE_BYTES"
    assert all(not u.image_ref and not u.image_refs for u in units)
    assert any(u.metadata["conversion_warnings"] for u in units)


def test_unavailable_source_cannot_establish_global_absence():
    spec = normalize_extraction_schema(
        {"type": "object", "properties": {"value": {"type": ["string", "null"]}}}
    )
    result = VerificationBatchResult(
        outcomes={
            "/value": VerificationOutcome(
                field_path="/value",
                status=VerificationStatus.SUPPORTED,
                raw_probabilities={"SUPPORTED": 0.99},
                model_id="decision",
            )
        }
    )
    payload = UnitInputPayload(
        unit=DocumentUnit(
            unit_id="u",
            document_id="d",
            ordinal=1,
            kind=UnitKind.DOCX_BLOCK,
            metadata={"unavailable_images": ["vector.emf"]},
        )
    )
    guarded = ExtractionPipelineService._guard_unavailable_sources(
        result, [(spec.fields[0], None, [])], {"u": payload}
    )
    assert guarded.outcomes["/value"].status == VerificationStatus.INSUFFICIENT_EVIDENCE
    assert guarded.outcomes["/value"].raw_probabilities == {"SUPPORTED": 0.99}
